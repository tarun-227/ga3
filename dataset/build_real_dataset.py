import json,ast,re,io,tokenize,pandas as pd,collections,difflib,hashlib
import openpyxl
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter
hdr=['sample_id','group_id','origin','case_type','label','task_description','db_api','n_lines','band','query_style','source','fix_type','code','fixed_code','vulnerable_lines','sink','n_tokens','source_url']
def band(n): return '1-5' if n<=5 else '6-10' if n<=10 else '11-20' if n<=20 else '21-40' if n<=40 else '41-100' if n<=100 else '>100'
def ntok(c):
    try: return sum(1 for t in tokenize.generate_tokens(io.StringIO(c).readline) if t.type not in (tokenize.NL,tokenize.NEWLINE,tokenize.INDENT,tokenize.DEDENT,tokenize.ENDMARKER))
    except Exception: return None
SINK2=re.compile(r'((?:\w+(?:\[[^\]]*\])?\.)*\w*(?:sql|query|command|exec|run|select|fetch|read_sql)\w*)\(',re.I)
SINK=re.compile(r'((?:\w+(?:\[[^\]]*\])?\.)*(?:execute|executemany|fetchrow|fetchval|fetchall|fetch|query|raw|run_query|exec|executescript|execute_sql|read_sql|read_sql_query))\(')
def db_api(c):
    t=c
    if re.search(r'sqlite',t,re.I): return 'sqlite3'
    if re.search(r'psycopg|asyncpg|\$1|RealDictCursor',t): return 'psycopg2' if 'psycopg' in t else 'other'
    if re.search(r'MySQLdb|mysql|pymysql|dictionary=True',t,re.I): return 'mysql'
    if re.search(r'\bsession\.|sqlalchemy|\btext\(|engine\.|db\.session',t): return 'sqlalchemy'
    if re.search(r'objects\.raw|django|connection\.cursor|\.extra\(',t): return 'django_orm'
    return None
def qstyle(lines):
    t=' '.join(lines)
    if re.search(r'\bf["\']',t): return 'f_string'
    if '.format(' in t: return 'format_method'
    if re.search(r'["\']\s*%\s*[\(\w]',t): return 'percent_format'
    if re.search(r'["\']\s*\+|\+\s*["\']|\+\s*str\(',t): return 'concatenation'
    return 'other'
def source(c):
    if re.search(r'request|\.args|\.form|\.GET|\.POST|cgi\.|FieldStorage|get_argument|self\.get\(',c): return 'web_request'
    if 'os.environ' in c: return 'environment'
    if re.search(r'\bopen\(|sys\.argv|csv\.',c): return 'file'
    return 'function_argument'
def fixtype(before,after):
    b=set(before.splitlines());a=[l for l in after.splitlines() if l not in b]
    added=' '.join(a)
    if re.search(r'\.(filter|query|objects)\b|\.delete\(\)|session\.query',added) : return 'orm_rewrite'
    if re.search(r'(\?|%s|:\w+|\$\d)',added) and re.search(r',\s*\(|,\s*\[|\bparams\b|values=',added): return 'parameterised_query'
    if re.search(r'\bint\(|\bfloat\(',added): return 'type_conversion'
    if re.search(r'\bnot in\b|\bin \(|ALLOWED|allowed|whitelist|\.get\(',added): return 'allow_list'
    return 'other'
def changed(before,after):
    bl=before.splitlines();al=after.splitlines();idx=[]
    for tag,i1,i2,j1,j2 in difflib.SequenceMatcher(None,bl,al).get_opcodes():
        if tag in('replace','delete'): idx+=list(range(i1+1,i2+1))
    return idx

def strip(src):
    try:
        toks=list(tokenize.generate_tokens(io.StringIO(src).readline))
    except Exception:
        toks=None
    lines=src.split(chr(10))
    comment_only=set()
    if toks:
        for t in toks:
            if t.type==tokenize.COMMENT:
                r,c=t.start
                if lines[r-1][:c].strip()=='' : comment_only.add(r-1)
                else: lines[r-1]=lines[r-1][:c].rstrip()
    else:
        for i,l in enumerate(lines):
            if l.strip().startswith('#'): comment_only.add(i)
            elif ' #' in l and l.count(chr(34))%2==0 and l.count(chr(39))%2==0: lines[i]=l.split(' #')[0].rstrip()
    doc=set()
    try:
        import textwrap
        t=ast.parse(textwrap.dedent(src))
        for n in ast.walk(t):
            if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef,ast.ClassDef)) and n.body:
                b=n.body[0]
                if isinstance(b,ast.Expr) and isinstance(getattr(b,'value',None),ast.Constant) and isinstance(b.value.value,str):
                    for k in range(b.lineno-1,b.end_lineno): doc.add(k)
    except Exception: pass
    out=[l.rstrip() for i,l in enumerate(lines) if i not in comment_only and i not in doc]
    res=chr(10).join(out).strip(chr(10))
    res=re.sub(chr(10)+'{3,}',chr(10)+chr(10),res)
    return res
rows=[];seen=set()
def add(group,origin,url,func_before,func_after):
    func_before=strip(func_before);func_after=strip(func_after)
    key=hashlib.md5(func_before.encode()).hexdigest()
    if key in seen: return
    seen.add(key)
    vl=changed(func_before,func_after);bl=func_before.splitlines()
    sink=None
    for i in vl:
        m=SINK.search(bl[i-1])
        if m: sink=m.group(1);break
    if not sink:
        m=SINK.search(func_before) or SINK2.search(func_before);sink=m.group(1) if m else None
    if not sink:
        vm=re.search(r'(\w+)\s*=',bl[vl[0]-1]) if vl else None
        if vm:
            cm=re.search(r'(\w+(?:\.\w+)*)\([^)]*\b'+re.escape(vm.group(1))+r'\b',func_before)
            sink=cm.group(1) if cm else None
    if not sink: sink='none'
    n=len(bl)
    rows.append(dict(group=group,origin=origin,url=url,before=func_before,after=func_after,vl=vl,sink=sink,
        db=db_api(func_before+func_after),qs=qstyle([bl[i-1] for i in vl]) ,src=source(func_before),ft=fixtype(func_before,func_after)))
# SVEN
for fn,split in (('sven89.jsonl','train'),('sven89_val.jsonl','val')):
    for l in open(fn):
        r=json.loads(l)
        sha=r['commit_link'].rstrip('/').split('/')[-1]
        add(sha[:12],f'real, SVEN {split}','https://'+r['commit_link'],r['func_src_before'],r['func_src_after'])
n_sven=len(rows)
# CrossVul (file-level -> changed functions)
def funcs(src):
    t=ast.parse(src);out={}
    lines=src.splitlines()
    for n in ast.walk(t):
        if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef)):
            start=min([n.lineno]+[d.lineno for d in n.decorator_list])
            out.setdefault(n.name,'\n'.join(lines[start-1:n.end_lineno]))
    return out
x=pd.read_pickle('py89.pkl');cv=0
for _,r in x.iterrows():
    try: fb,fa=funcs(r.vulnerable_code),funcs(r.fixed_code)
    except SyntaxError: continue
    for name,src in fb.items():
        if name in fa and fa[name]!=src and re.search(r'execute|SELECT|INSERT|UPDATE|DELETE|query',src,re.I) and len(src.splitlines())<=100:
            add('crossvul-'+r.file_pair_id,'real, CrossVul','https://zenodo.org/records/4734050',src,fa[name]);cv+=1
print('sven',n_sven,'crossvul',cv)
# CVEfixes: try function extraction
cf=pd.read_pickle('cf89.pkl');cfn=0
for _,r in cf.iterrows():
    try: fb,fa=funcs(r.vulnerable_code),funcs(r.fixed_code)
    except Exception: continue
    for name,src in fb.items():
        if name in fa and fa[name]!=src:
            add(r.cve_id,'real, CVEfixes',r.repo_url+'/commit/'+r.hash,src,fa[name]);cfn+=1
print('cvefixes funcs',cfn)

def human(name): return re.sub(r'(?<=[a-z0-9])(?=[A-Z])',' ',name).replace('_',' ').strip().lower()
def describe(r):
    code=r['before'];bl=code.splitlines()
    m=re.search(r'def\s+(\w+)\s*\(([^)]*)\)',code)
    name=m.group(1) if m else 'function';args=[a.strip().split('=')[0].strip() for a in (m.group(2) if m else '').split(',') if a.strip()]
    args=[a for a in args if a not in('self','cls','conn','cursor','db','cur','connection','c','dbh','con','curs') and not a.startswith('*')]
    sql=' '.join(bl[i-1] for i in r['vl']) or code
    up=sql.upper()
    stmt=next((k for k in('INSERT','UPDATE','DELETE','SELECT') if k in up),None)
    if stmt is None:
        up2=code.upper();stmt=next((k for k in('INSERT','UPDATE','DELETE','SELECT') if k in up2),None)
    full=code
    tm=re.search(r'(?:FROM|INTO|UPDATE)\s+([A-Za-z_][\w.]*)',full,re.I)
    table=tm.group(1) if tm and tm.group(1).lower() not in('set','select') else None
    cm=re.findall(r'(\w+)\s*(?:=|LIKE|IN|<|>)\s*(?:[\'"%?\(\{]|%s|\w*\s*\+)',sql,re.I)
    cols=[c for c in dict.fromkeys(cm) if c.lower() not in('where','and','or','set','values','from','not','is')][:3]
    used_args=[a for a in args if re.search(r'\b'+re.escape(a)+r'\b',sql)] or [a for a in args if re.search(r'\b'+re.escape(a)+r'\b',code.split(chr(10),1)[-1])]
    argtxt=(' and '.join(used_args[:2])) if used_args else 'the given values'
    pre='Function '+human(name)+' should '
    if stmt=='SELECT':
        d=f"look up rows{' in the '+table+' table' if table else ''}{' matching '+' and '.join(cols) if cols else ' using '+argtxt}, and return the results"
    elif stmt=='INSERT':
        d=f"store a new record{' in the '+table+' table' if table else ''} built from {argtxt}"
    elif stmt=='UPDATE':
        d=f"change existing records{' in the '+table+' table' if table else ''}{' selected by '+' and '.join(cols) if cols else ' using '+argtxt}"
    elif stmt=='DELETE':
        d=f"remove records{' from the '+table+' table' if table else ''}{' selected by '+' and '.join(cols) if cols else ' using '+argtxt}"
    else:
        d=f"run a database operation{' on the '+table+' table' if table else ''} based on {argtxt}"
    return pre+d+'.'
def guess_db(r):
    if r['db']: return r['db']
    t=r['before']+chr(10)+r['after']
    if re.search(r'ILIKE|RETURNING|::\w+|\bserial\b|psycopg',t,re.I): return 'psycopg2'
    if re.search(r'\?',r['after']) and not re.search(r'%s',r['after']): return 'sqlite3'
    if re.search(r'%s',r['after']):
        return 'mysql' if re.search(r'mysql|MySQL|db\.cursor|cnx|DictCursor',t) else 'psycopg2'
    if re.search(r'sqlite|\.execute\(.*\?',t): return 'sqlite3'
    return 'other'
used=collections.Counter()
for r in rows:
    r['desc']=describe(r);r['db_filled']=r['db'] is None;r['db']=guess_db(r)
    used[r['desc']]+=1
# disambiguate duplicates with the commit group id
for r in rows:
    if used[r['desc']]>1: r['desc']=r['desc'][:-1]+' (commit '+r['group'][:7]+').'

out=[];i=0
def mk(i,r,vuln):
    code=r['before'] if vuln else r['after']
    lines=code.splitlines()
    return [f'real_{i:04d}',r['group'],r['origin'],'real_vulnerable' if vuln else 'real_fixed',1 if vuln else 0,r['desc'],r['db'],len(lines),band(len(lines)),
        r['qs'] if vuln else None,r['src'] if vuln else None,r['ft'] if vuln else None,code,r['after'] if vuln else None,str(r['vl']) if vuln else '[]',r['sink'],ntok(code),r['url']]
for r in rows:
    out.append(mk(i,r,True));i+=1
    out.append(mk(i,r,False));i+=1
wb=openpyxl.Workbook();ws=wb.active;ws.title='real_rows'
ws.append(hdr)
for o in out: ws.append(o)
for c in ws[1]: c.font=Font(bold=True)
for k,wd in enumerate([11,14,20,16,6,18,12,8,8,16,18,20,70,70,12,18,9,50],1): ws.column_dimensions[get_column_letter(k)].width=wd
ws.freeze_panes='A2'
# highlight blanks
blank=PatternFill('solid',fgColor='FFF2CC')
bl=wb.create_sheet('blank_columns')
bl.append(['column','rows','blank','note'])
for ci,h in enumerate(hdr):
    nb=sum(1 for o in out if o[ci] is None or o[ci]=='')
    bl.append([h,len(out),nb,''])
wb.save('/home/user/ga3/dataset/sqli_real_dataset.xlsx')
print(len(out),'rows');
for ci,h in enumerate(hdr): print(h,sum(1 for o in out if o[ci] is None or o[ci]==''))
print(collections.Counter(o[6] for o in out if o[4]==1),collections.Counter(o[8] for o in out if o[4]==1),collections.Counter(o[9] for o in out if o[4]==1),collections.Counter(o[11] for o in out if o[4]==1))
