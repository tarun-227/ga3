import openpyxl,re,difflib,collections
P='dataset/sqli_real_dataset.xlsx'
wb=openpyxl.load_workbook(P);ws=wb['real_rows']
rows=[list(r) for r in ws.iter_rows(min_row=2,values_only=True)]
H=[c.value for c in ws[1]];ix={h:i for i,h in enumerate(H)}
def diff(a,b):
    al,bl=a.splitlines(),b.splitlines();rem=[];add=[]
    for tag,i1,i2,j1,j2 in difflib.SequenceMatcher(None,al,bl).get_opcodes():
        if tag!='equal': rem+=al[i1:i2];add+=bl[j1:j2]
    return rem,add
def qstyle(rem):
    t=' '.join(rem)
    if re.search(r'\bf["\']',t): return 'f_string'
    if '.format(' in t: return 'format_method'
    if re.search(r'["\']\s*%\s*[\(\w]|\bquery\s*%|%\s*\(',t) or re.search(r"%['\"]?\s*$",t): return 'percent_format'
    if re.search(r'["\']\s*\+|\+\s*["\']|\+\s*str\(',t): return 'concatenation'
    return 'other'
def source(code):
    if re.search(r'\brequest\.|flask|cgi\.|FieldStorage|get_argument|form_dict|bottle|\.GET\b|\.POST\b|self\.request|web\.input|params\[|request\[|query_params|\.args\b|whitelist',code): return 'web_request'
    if re.search(r'message\.|trigger|\bbot\b|chat_id|update\.',code): return 'other'
    return 'function_argument'
def dbapi(before,after):
    t=before+'\n'+after
    if re.search(r'frappe\.db|\bdb\.exec\(|self\.db\.exec|execute_query|db\.add\(|bot\.db',t): return 'other'
    if re.search(r'sqlalchemy|\btext\(|db\.engine|db\.session|session\.execute',t): return 'sqlalchemy'
    if re.search(r'\$\d',after): return 'other'
    if re.search(r'objects\.raw|connection\.cursor\(\)|\.extra\(',t): return 'django_orm'
    if re.search(r'mysql|MySQL|pymysql|cnx|dictionary=True|DictCursor|MySQLdb',t): return 'mysql'
    if re.search(r'psycopg|ILIKE|RETURNING|::\w+|to_tsquery|\bserial\b|array_append|unnest',t,re.I): return 'psycopg2'
    if re.search(r'sqlite',t,re.I): return 'sqlite3'
    if re.search(r'\?',after) and not re.search(r'%s|%\(',after): return 'sqlite3'
    if re.search(r':\w+',after) and re.search(r'cursor\.execute',after): return 'sqlite3'
    if re.search(r'%\(\w+\)s|%s',after): return 'psycopg2'
    return 'other'
cnt=collections.Counter()
for k in range(0,len(rows),2):
    v,f=rows[k],rows[k+1]
    rem,add=diff(v[ix['code']],f[ix['code']])
    qs=qstyle(rem);src=source(v[ix['code']]);api=dbapi(v[ix['code']],f[ix['code']])
    ft='parameterised_query'
    for r in (v,f):
        r[ix['db_api']]=api;r[ix['source']]=src;r[ix['fix_type']]=ft
    v[ix['query_style']]=qs
    f[ix['query_style']]='parameterised'
    f[ix['fixed_code']]=f[ix['code']]
    cnt[(api,qs,src)]+=1
ws.delete_rows(2,ws.max_row)
for r in rows: ws.append(r)
bs=wb['blank_columns'];bs.delete_rows(2,bs.max_row)
for h in H:
    nb=sum(1 for r in rows if r[ix[h]] in (None,''))
    bs.append([h,len(rows),nb,''])
wb.save(P)
for name in ('db_api','query_style','source','fix_type'): print(name,collections.Counter(r[ix[name]] for r in rows if r[ix['label']]==1))
print({h:sum(1 for r in rows if r[ix[h]] in (None,'')) for h in H})
